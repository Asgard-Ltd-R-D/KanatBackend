import csv
import io
import re
import urllib.parse
from datetime import datetime, timezone

import openpyxl
import weasyprint
from jinja2 import Environment
from openpyxl.styles import Alignment, Font, PatternFill

from app.models import Session
from app.services import groups

_REPORT_HTML = """<!DOCTYPE html>
<html>
<head>
  <meta charset="utf-8">
  <style>
    * { box-sizing: border-box; margin: 0; padding: 0; }
    body { font-family: -apple-system, Inter, system-ui, sans-serif; color: #1e293b; padding: 40px; font-size: 13px; }
    h1 { font-size: 20px; font-weight: 700; margin-bottom: 4px; }
    .subtitle { color: #64748b; font-size: 12px; margin-bottom: 28px; }
    .summary { display: flex; gap: 16px; margin-bottom: 32px; }
    .card { border: 1px solid #e2e8f0; border-radius: 8px; padding: 14px 20px; background: #f8fafc; }
    .card-label { font-size: 10px; color: #94a3b8; text-transform: uppercase; letter-spacing: 0.06em; margin-bottom: 4px; }
    .card-value { font-size: 24px; font-weight: 700; color: #0f172a; }
    h2 { font-size: 14px; font-weight: 700; margin-bottom: 10px; }
    .groups { margin-bottom: 32px; }
    .groups th, .groups td { padding: 8px; }
    .groups td:first-child { white-space: nowrap; }
    .note { color: #64748b; font-size: 11px; margin-top: 8px; }
    table { width: 100%; border-collapse: collapse; }
    thead tr { background: #1e293b; }
    th { text-align: left; padding: 9px 14px; color: #fff; font-weight: 600; font-size: 11px; text-transform: uppercase; letter-spacing: 0.05em; }
    td { padding: 8px 14px; border-bottom: 1px solid #f1f5f9; }
    tr:nth-child(even) td { background: #f8fafc; }
    .footer { margin-top: 28px; font-size: 10px; color: #94a3b8; text-align: right; }
  </style>
</head>
<body>
  <h1>{{ session.name }}</h1>
  <p class="subtitle">
    {{ session.started_at.strftime('%Y-%m-%d %H:%M') }} UTC
    &nbsp;&middot;&nbsp; Duration: {{ duration }}
    &nbsp;&middot;&nbsp; {{ session.status | capitalize }}
  </p>
  <div class="summary">
    <div class="card">
      <div class="card-label">Total Bullets</div>
      <div class="card-value">{{ bullets | length }}</div>
    </div>
  </div>
  <div class="groups">
    <h2>Groups</h2>
    {% if grouping.groups %}
    <table>
      <thead>
        <tr>
          <th>Target</th>
          <th>Bullet Holes</th>
          <th>MPI X (mm)</th>
          <th>MPI Y (mm)</th>
          <th>CEP (mm)</th>
          <th>Mean Radius (mm)</th>
          <th>RMS Radius (mm)</th>
          <th>Extreme Spread (mm)</th>
        </tr>
      </thead>
      <tbody>
        {% for g in grouping.groups %}
        <tr>
          <td>Target {{ g.target_index + 1 }}</td>
          <td>{{ g.n }}</td>
          <td>{{ g.mpi_x_mm | mm }}</td>
          <td>{{ g.mpi_y_mm | mm }}</td>
          <td>{{ g.cep_mm | mm }}</td>
          <td>{{ g.mean_radius_mm | mm }}</td>
          <td>{{ g.rms_radius_mm | mm }}</td>
          <td>{{ g.extreme_spread_mm | mm }}</td>
        </tr>
        {% endfor %}
      </tbody>
    </table>
    {% endif %}
    <p class="note">
      In no Group: Misses: {{ grouping.misses }}
      &nbsp;&middot;&nbsp; Without a Position: {{ grouping.without_position }}
    </p>
  </div>
  <table>
    <thead>
      <tr>
        <th>#</th>
        <th>Detected At (UTC)</th>
        <th>Position X</th>
        <th>Position Y</th>
      </tr>
    </thead>
    <tbody>
      {% for b in bullets %}
      <tr>
        <td>{{ b.rank }}</td>
        <td>{{ b.first_seen_at.strftime('%H:%M:%S') }}</td>
        <td>{{ "%.1f" | format(b.position_x) }}</td>
        <td>{{ "%.1f" | format(b.position_y) }}</td>
      </tr>
      {% endfor %}
    </tbody>
  </table>
  <div class="footer">Generated {{ generated_at }}</div>
</body>
</html>"""


def _session_duration(session: Session) -> str:
    if session.ended_at and session.started_at:
        delta = session.ended_at - session.started_at
        total_s = int(delta.total_seconds())
        m, s = divmod(total_s, 60)
        return f"{m}m {s:02d}s"
    return "—"


def _ascii_slug(name: str) -> str:
    slug = name.encode("ascii", errors="replace").decode()
    slug = re.sub(r"[^\w\s-]", "_", slug)
    slug = re.sub(r"[\s]+", "_", slug)
    return slug.strip("_") or "session"


def _content_disposition(session: Session, ext: str) -> str:
    date_str = session.started_at.strftime("%Y%m%d")
    fallback = f"session_{_ascii_slug(session.name)}_{date_str}.{ext}"
    utf8_name = f"session_{session.name.replace(' ', '_')}_{date_str}.{ext}"
    encoded = urllib.parse.quote(utf8_name, safe="-._~")
    return f"attachment; filename=\"{fallback}\"; filename*=UTF-8''{encoded}"


_jinja_env = Environment(autoescape=True)
# A Group's measure to 0.1 mm; a Group of one has no spread
_jinja_env.filters["mm"] = lambda v: "—" if v is None else f"{v:.1f}"


def generate_pdf(session: Session) -> bytes:
    bullets = sorted(session.bullet_holes, key=lambda b: b.rank)
    html = _jinja_env.from_string(_REPORT_HTML).render(
        session=session,
        bullets=bullets,
        grouping=groups.grouping(bullets),
        duration=_session_duration(session),
        generated_at=datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
    )
    return weasyprint.HTML(string=html, base_url=None).write_pdf()


def generate_csv(session: Session) -> str:
    bullets = sorted(session.bullet_holes, key=lambda b: b.rank)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["rank", "first_seen_at", "position_x", "position_y", "source"])
    for b in bullets:
        writer.writerow([
            b.rank,
            b.first_seen_at.isoformat(),
            b.position_x,
            b.position_y,
            b.source,
        ])
    return output.getvalue()


def _append_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.fill = PatternFill("solid", fgColor="1E293B")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")


def _fit_columns(ws) -> None:
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 2, 12)


def generate_excel(session: Session) -> bytes:
    bullets = sorted(session.bullet_holes, key=lambda b: b.rank)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bullet Holes"
    _append_header(ws, ["#", "Detected At (UTC)", "Position X", "Position Y", "Source"])

    for b in bullets:
        ws.append([
            b.rank,
            b.first_seen_at.strftime("%Y-%m-%d %H:%M:%S"),
            b.position_x,
            b.position_y,
            b.source.capitalize(),   # Model or Manual, as the app shows it
        ])
    _fit_columns(ws)

    # The same rows and counts as the detail response's grouping
    grouping = groups.grouping(bullets)
    ws = wb.create_sheet("Groups")
    _append_header(ws, ["Target", "Bullet Holes", "MPI X (mm)", "MPI Y (mm)", "CEP (mm)",
                        "Mean Radius (mm)", "RMS Radius (mm)", "Extreme Spread (mm)"])
    for g in grouping.groups:
        ws.append([f"Target {g.target_index + 1}", g.n, g.mpi_x_mm, g.mpi_y_mm, g.cep_mm,
                   g.mean_radius_mm, g.rms_radius_mm, g.extreme_spread_mm])
        for cell in ws[ws.max_row][2:]:
            cell.number_format = "0.0"
    ws.append([])
    ws.append(["Misses", grouping.misses])
    ws.append(["Without a Position", grouping.without_position])
    _fit_columns(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
